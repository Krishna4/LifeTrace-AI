import os
import logging
from typing import Any, Dict
from src.backend.utils.memory_monitor import enforce_memory_ceiling

logger = logging.getLogger(__name__)


def extract_docx_text(docx_path: str) -> Dict[str, Any]:
    """
    Extracts clean text from Microsoft Word (.docx) documents.
    Uses python-docx if available, or built-in zipfile + ElementTree XML parsing.
    """
    enforce_memory_ceiling()

    if not os.path.exists(docx_path):
        return {
            "status": "PARTIAL_SUCCESS",
            "text": f"DOCX File: {os.path.basename(docx_path)}",
            "error": f"File does not exist: {docx_path}",
        }

    try:
        # 1. Try python-docx if available in environment
        try:
            import docx
            doc = docx.Document(docx_path)
            full_text = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        full_text.append(row_text)
            text = "\n".join(full_text)
        except Exception:
            # 2. Fallback to standard library zipfile + ElementTree XML parsing
            import zipfile
            import xml.etree.ElementTree as ET

            with zipfile.ZipFile(docx_path, "r") as z:
                xml_content = z.read("word/document.xml")
                tree = ET.fromstring(xml_content)
                texts = []
                for elem in tree.iter():
                    if elem.tag.endswith("}t") and elem.text:
                        texts.append(elem.text)
                text = " ".join(texts)

        if not text.strip():
            text = f"DOCX Document: {os.path.basename(docx_path)}"

        return {
            "status": "COMPLETED",
            "text": text.strip(),
            "error": None,
        }
    except Exception as e:
        logger.error(f"DOCX extraction error for {docx_path}: {e}")
        return {
            "status": "PARTIAL_SUCCESS",
            "text": f"DOCX content from {os.path.basename(docx_path)}",
            "error": str(e),
        }
